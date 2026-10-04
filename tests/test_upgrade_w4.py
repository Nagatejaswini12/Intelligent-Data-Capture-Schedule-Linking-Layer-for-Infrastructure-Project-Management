"""Upgrade W4: Primavera P6 XER schedule import (round trip against the synthetic CSV schedule)."""
from __future__ import annotations

from pathlib import Path

import pytest

from p2e.plan import importers as imp

SCHEDULE = Path(__file__).resolve().parents[1] / "data" / "synthetic" / "schedule" / "schedule.csv"


def to_xer(rows: list[dict]) -> bytes:
    """Minimal P6-style XER for the synthetic schedule: WBS, tasks, predecessors, Discipline/Area/Activity Type codes."""
    wbs = [r for r in rows if r["node_type"] != "activity"]
    wid = {r["node_id"]: str(100 + i) for i, r in enumerate(wbs)}
    tid = {r["node_id"]: str(5000 + i) for i, r in enumerate(r for r in rows if r["node_type"] == "activity")}
    acts = [r for r in rows if r["node_type"] == "activity"]
    types = {"Discipline": "discipline", "Area": "area", "Activity Type": "activity_type"}
    values = sorted({(t, a[f]) for a in acts for t, f in types.items() if a[f]})
    vid = {v: str(900 + i) for i, v in enumerate(values)}
    tyid = {t: str(80 + i) for i, t in enumerate(types)}
    out = ["ERMHDR\t19.12\t2026-09-16\tProject\tadmin\tP2E\tdbxDatabaseNoName\tProject Management\tINR"]

    def table(name, fields, recs):
        out.append(f"%T\t{name}")
        out.append("%F\t" + "\t".join(fields))
        out.extend("%R\t" + "\t".join(str(x) for x in rec) for rec in recs)

    table("PROJECT", ["proj_id", "proj_short_name", "last_recalc_date"], [["1", "CGS-EXP-01", "2026-09-16 00:00"]])
    table("PROJWBS", ["wbs_id", "proj_id", "parent_wbs_id", "proj_node_flag", "wbs_short_name", "wbs_name"],
          [[wid[r["node_id"]], "1", wid.get(r["parent_id"], ""), "Y" if not r["parent_id"] else "N",
            r["node_id"] if r["node_type"] == "summary" or not r["parent_id"] else r["node_id"][len(r["parent_id"]) + 1:], r["name"]]
           for r in wbs])
    table("TASK", ["task_id", "proj_id", "wbs_id", "task_code", "task_name", "task_type", "target_start_date", "target_end_date",
                   "act_start_date", "act_end_date"],
          [[tid[a["node_id"]], "1", wid[a["parent_id"]], a["node_id"], a["name"], "TT_Task", f"{a['planned_start']} 08:00",
            f"{a['planned_finish']} 17:00", f"{a['actual_start']} 08:00" if a["actual_start"] else "",
            f"{a['actual_finish']} 17:00" if a["actual_finish"] else ""] for a in acts])
    table("TASKPRED", ["task_pred_id", "task_id", "pred_task_id", "pred_type", "lag_hr_cnt"],
          [[str(i), tid[a["node_id"]], tid[p.split(":")[0]], "PR_" + p.split(":")[1][:2], int(p.split(":")[1][2:]) * 8]
           for i, (a, p) in enumerate((a, p) for a in acts for p in filter(None, a["predecessors"].split(";")))])
    table("ACTVTYPE", ["actv_code_type_id", "actv_code_type"], [[tyid[t], t] for t in types])
    table("ACTVCODE", ["actv_code_id", "actv_code_type_id", "short_name"], [[vid[v], tyid[v[0]], v[1]] for v in values])
    table("TASKACTV", ["task_id", "actv_code_type_id", "actv_code_id"],
          [[tid[a["node_id"]], tyid[t], vid[(t, a[f])]] for a in acts for t, f in types.items() if a[f]])
    out.append("%E")
    return ("\r\n".join(out) + "\r\n").encode("cp1252", errors="replace")


@pytest.fixture(scope="module")
def csv_rows():
    return imp.read_schedule(SCHEDULE).rows


def test_xer_round_trip_matches_csv(csv_rows, tmp_path):
    path = tmp_path / "schedule.xer"
    path.write_bytes(to_xer(csv_rows))
    parsed = imp.read_schedule(path)
    assert parsed.format == "xer" and str(parsed.data_date) == "2026-09-16"
    assert imp.compare_schedules(csv_rows, parsed.rows) == []
    acts = {r["node_id"]: r for r in parsed.rows if r["node_type"] == "activity"}
    csv_acts = {r["node_id"]: r for r in csv_rows if r["node_type"] == "activity"}
    assert all(acts[k]["activity_type"] == v["activity_type"] for k, v in csv_acts.items())
    nodes = imp.validate_rows(parsed.rows)                         # the normal Phase 1 validation accepts it
    assert sum(n.node_type == "activity" for n in nodes) == 317


@pytest.mark.parametrize("raw,msg", [
    (b"not an xer", "missing ERMHDR"),
    (b"ERMHDR\t19.12\r\n%T\tPROJECT\r\n%F\tproj_id\r\n%R\t1\r\n%E\r\n", "no PROJWBS, TASK table"),
])
def test_xer_rejects_bad_files(raw, msg):
    with pytest.raises(imp.ScheduleValidationError, match=msg):
        imp.parse_xer(raw)


# ----------------------------------------------------------------------------- .docx / .csv field reports

import csv as _csv  # noqa: E402
import io  # noqa: E402
import zipfile  # noqa: E402
from datetime import datetime  # noqa: E402
from xml.sax.saxutils import escape  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from p2e.extract.pipeline import extract_bytes, load_project_vocab  # noqa: E402
from tests.test_phase2 import H, REPORTS, SHEETS, api  # noqa: E402,F401
from tests.test_phase3 import GLOSSARY  # noqa: E402

PV = load_project_vocab(GLOSSARY)
FIELDS = ("activity_text", "event_type", "event_date", "quantity", "unit", "discipline", "area", "tags", "source_text")


def to_docx(text: str) -> bytes:
    """Minimal Word document: one paragraph per line (what a supervisor gets by pasting the DPR into Word)."""
    body = "".join(f'<w:p><w:r><w:t xml:space="preserve">{escape(line)}</w:t></w:r></w:p>' for line in text.splitlines())
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   f"<w:body>{body}</w:body></w:document>")
    return buf.getvalue()


def to_csv(xlsx: bytes) -> bytes:
    ws = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True).worksheets[0]
    out = io.StringIO()
    w = _csv.writer(out)
    for row in ws.iter_rows(values_only=True):
        w.writerow(["" if v is None else v.date().isoformat() if isinstance(v, datetime) else v for v in row])
    return out.getvalue().encode("utf-8")


def same(a, b):
    return [tuple(getattr(i, f) for f in FIELDS) for i in a.items] == [tuple(getattr(i, f) for f in FIELDS) for i in b.items]


@pytest.mark.parametrize("path", REPORTS[:6], ids=lambda p: p.name)
def test_docx_report_extracts_like_the_text(path):
    text = path.read_bytes()
    txt, docx = extract_bytes(text, ".txt", PV), extract_bytes(to_docx(text.decode("utf-8")), ".docx", PV)
    assert txt.items and same(txt, docx) and docx.report_date == txt.report_date and not docx.issues


def test_csv_sheet_extracts_like_the_workbook():
    """Single-sheet trackers: the CSV export gives the same events as the workbook."""
    checked = 0
    for path in SHEETS:
        if len(load_workbook(path, read_only=True).worksheets) != 1:
            continue
        xl = extract_bytes(path.read_bytes(), ".xlsx", PV)
        cs = extract_bytes(to_csv(path.read_bytes()), ".csv", PV)
        assert xl.items and [tuple(getattr(i, f) for f in FIELDS[:-1]) for i in xl.items] == \
               [tuple(getattr(i, f) for f in FIELDS[:-1]) for i in cs.items], path.name
        checked += 1
    assert checked >= 1


def test_api_accepts_docx_and_csv_with_evidence(api):
    c, base = api
    rep = REPORTS[3]
    d = c.post(f"{base}/documents", files={"file": (rep.stem + ".docx", to_docx(rep.read_text(encoding="utf-8")))}, headers=H)
    assert d.status_code == 201 and d.json()["kind"] == "dpr_text", d.text
    p = c.post(f"{base}/documents/{d.json()['id']}/process", headers=H).json()
    assert p["outcome"] == "processed" and p["run"]["events_total"] > 0 and p["run"]["events_invalid"] == 0
    e = c.get(f"{base}/events?document_id={d.json()['id']}", headers=H).json()["items"][0]
    ev = c.get(f"{base}/events/{e['id']}/evidence", headers=H).json()
    assert ev["found_in_source"]
    sheet = next(s for s in SHEETS if len(load_workbook(s, read_only=True).worksheets) == 1)
    s = c.post(f"{base}/documents", files={"file": (sheet.stem + ".csv", to_csv(sheet.read_bytes()))}, headers=H)
    assert s.status_code == 201 and s.json()["kind"] == "spreadsheet", s.text
    p = c.post(f"{base}/documents/{s.json()['id']}/process", headers=H).json()
    assert p["outcome"] == "processed" and p["run"]["events_total"] > 0
    e = c.get(f"{base}/events?document_id={s.json()['id']}", headers=H).json()["items"][0]
    assert c.get(f"{base}/events/{e['id']}/evidence", headers=H).json()["found_in_source"]
    for name, data, code in (("macro.docx", b"PK\x03\x04junk", 422), ("x.docm", b"PK", 415), ("bad.csv", b"\x00\x01", 422)):
        assert c.post(f"{base}/documents", files={"file": (name, data)}, headers=H).status_code == code, name
